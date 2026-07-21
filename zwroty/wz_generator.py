"""Генерація документів WZ на боці сервера.

Раніше цю логіку виконував Google Apps Script (заповнення Google-Sheets-шаблону
та експорт у PDF). Тепер усе робиться тут:

* для кожного замовлення — окремий ``.xlsx`` файл;
* один спільний ``.pdf`` з усіма аркушами, де кожен аркуш повторюється
  ``COPIES_PER_SHEET`` разів підряд (3 копії).

Файли складаються у тимчасову теку ``wz_tmp/<batch_id>/`` разом із
``manifest.json``. Apps Script завантажує їх за токенами і після успіху
викликає confirm-ендпоінт, який видаляє теку.

Мапінг клітинок шаблону повторює ``fillSheet`` зі старого Apps Script:
    A = L.p.,  B = referencja handlowa (sku_hand),  C = referencja logistyczna
    (sku_log),  D = nazwa towaru (descript),  E = ilość (qty),  J = uwagi (reasone).
Дані займають рядки 7..16 (до 10 позицій на аркуш).
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import uuid
from datetime import datetime

from django.conf import settings
from openpyxl import load_workbook

# --- Конфіг ------------------------------------------------------------------

TEMPLATE_PATH = os.path.join(settings.BASE_DIR, "wz_exsamle.xlsx")
WZ_TMP_ROOT = os.path.join(settings.BASE_DIR, "wz_tmp")
TEMPLATE_SHEET_NAME = "WZ"

COPIES_PER_SHEET = 3     # скільки разів повторити кожен аркуш у спільному PDF
LINES_PER_SHEET = 10     # позицій (рядків) на один аркуш

# Бінарник LibreOffice для конвертації xlsx -> pdf.
# Можна перевизначити в settings.py: SOFFICE_BIN = "/usr/bin/soffice"
SOFFICE_BIN = getattr(settings, "SOFFICE_BIN", "soffice")


# --- Допоміжні ---------------------------------------------------------------

def _pad_wz(wz) -> str:
    return str(wz).zfill(4)


def _chunks(seq, size):
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


def _safe_filename(name: str) -> str:
    """Прибираємо символи, неприпустимі у назвах файлів (зберігаючи пробіли/UTF)."""
    return re.sub(r'[\\/:*?"<>|]', "_", name).strip()


def _order_filename(order) -> str:
    year = datetime.now().year
    base = f"BWWZ{_pad_wz(order['wz_nr'])}.{year} zwrot {order['shop_nr']}"
    return _safe_filename(base) + ".xlsx"


def _apply_page_setup(ws):
    """copy_worksheet не переносить параметри друку — виставляємо вручну."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:M21"


def _fill_sheet(ws, order, lines, index):
    """Заповнює один аркуш даними замовлення та порцією позицій ``lines``.

    ``index`` — номер порції (0,1,2...), потрібен для наскрізної нумерації L.p.
    """
    # Шапка
    ws["C3"] = order["ship_doc"]
    ws["E3"] = order["shop_desct"]
    ws["J2"] = order["shop_nr"]
    ws["L5"] = _pad_wz(order["wz_nr"])
    ws["H7"] = order["type_of_delivery"]
    ws["F7"] = order["bw_nr"]
    ws["B20"] = order["data"]
    ws["B19"] = order["user_name"]

    # Очищаємо область даних (на випадок копії з непорожнього аркуша)
    for r in range(7, 7 + LINES_PER_SHEET):
        for col in ("A", "B", "C", "D", "E", "J"):
            ws[f"{col}{r}"] = None

    # Позиції
    for i, line in enumerate(lines):
        r = 7 + i
        ws[f"A{r}"] = i + 1 + index * LINES_PER_SHEET
        ws[f"B{r}"] = line.get("sku_hand", "")
        ws[f"C{r}"] = line.get("sku_log", "")
        ws[f"D{r}"] = line.get("descript", "")
        qty = line.get("qty", "")
        try:
            ws[f"E{r}"] = int(qty)
        except (TypeError, ValueError):
            ws[f"E{r}"] = qty
        ws[f"J{r}"] = line.get("reasone", "")


# --- Побудова xlsx -----------------------------------------------------------

def build_order_workbook(order):
    """Один ``.xlsx`` для одного замовлення (кілька аркушів, якщо > 10 позицій)."""
    wb = load_workbook(TEMPLATE_PATH)
    template_ws = wb[TEMPLATE_SHEET_NAME]

    chunks = list(_chunks(order["lines"], LINES_PER_SHEET)) or [[]]
    for chunk_index, chunk in enumerate(chunks):
        if chunk_index == 0:
            ws = template_ws
            ws.title = "WZ"
        else:
            ws = wb.copy_worksheet(template_ws)
            ws.title = f"WZ_{chunk_index + 1}"
        _apply_page_setup(ws)
        _fill_sheet(ws, order, chunk, chunk_index)
    return wb


def build_combined_workbook(orders):
    """Один робочий зошит з усіма аркушами; кожен аркуш ×COPIES_PER_SHEET підряд."""
    wb = load_workbook(TEMPLATE_PATH)
    template_ws = wb[TEMPLATE_SHEET_NAME]

    sheet_no = 0
    for order in orders:
        chunks = list(_chunks(order["lines"], LINES_PER_SHEET)) or [[]]
        for chunk_index, chunk in enumerate(chunks):
            for _copy in range(COPIES_PER_SHEET):
                sheet_no += 1
                ws = wb.copy_worksheet(template_ws)
                ws.title = f"WZ_{sheet_no}"
                _apply_page_setup(ws)
                _fill_sheet(ws, order, chunk, chunk_index)

    wb.remove(template_ws)  # прибираємо порожній шаблонний аркуш
    return wb


# --- Конвертація в PDF -------------------------------------------------------

def convert_xlsx_to_pdf(xlsx_path, out_dir):
    """Конвертує xlsx у pdf через LibreOffice headless. Повертає шлях до pdf."""
    # Окремий профіль на кожен виклик — інакше паралельні soffice конфліктують.
    profile = f"file:///tmp/lo_profile_{uuid.uuid4().hex}"
    subprocess.run(
        [
            SOFFICE_BIN,
            "-env:UserInstallation=" + profile,
            "--headless",
            "--convert-to",
            "pdf:calc_pdf_Export",
            "--outdir",
            out_dir,
            xlsx_path,
        ],
        check=True,
        timeout=180,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    pdf_name = os.path.splitext(os.path.basename(xlsx_path))[0] + ".pdf"
    return os.path.join(out_dir, pdf_name)


# --- Оркестрація batch -------------------------------------------------------

def generate_wz_batch(orders):
    """Генерує всі файли й повертає ``(batch_id, files)``.

    ``files`` — список ``{"token", "filename", "path", "type"}``; окрім того
    у теці зберігається ``manifest.json``.
    """
    batch_id = uuid.uuid4().hex
    batch_dir = os.path.join(WZ_TMP_ROOT, batch_id)
    os.makedirs(batch_dir, exist_ok=True)

    files = []

    # 1) Окремі xlsx на кожне замовлення
    for order in orders:
        filename = _order_filename(order)
        path = os.path.join(batch_dir, filename)
        build_order_workbook(order).save(path)
        files.append(
            {"token": uuid.uuid4().hex, "filename": filename, "path": path, "type": "xlsx"}
        )

    # 2) Спільний PDF (3 копії кожного аркуша)
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    combined_xlsx = os.path.join(batch_dir, f"_combined_{ts}.xlsx")
    build_combined_workbook(orders).save(combined_xlsx)
    pdf_path = convert_xlsx_to_pdf(combined_xlsx, batch_dir)
    pdf_name = f"BWWZ_COMBINED_{ts}.pdf"
    final_pdf = os.path.join(batch_dir, pdf_name)
    if os.path.abspath(pdf_path) != os.path.abspath(final_pdf):
        os.replace(pdf_path, final_pdf)
    os.remove(combined_xlsx)  # проміжний файл не віддаємо
    files.append(
        {"token": uuid.uuid4().hex, "filename": pdf_name, "path": final_pdf, "type": "pdf"}
    )

    manifest = {"batch_id": batch_id, "created": datetime.now().isoformat(), "files": files}
    with open(os.path.join(batch_dir, "manifest.json"), "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=2)

    return batch_id, files


def _batch_dir(batch_id):
    # batch_id — hex uuid; захист від path traversal
    if not re.fullmatch(r"[0-9a-f]{32}", batch_id or ""):
        return None
    path = os.path.join(WZ_TMP_ROOT, batch_id)
    return path if os.path.isdir(path) else None


def resolve_file(batch_id, token):
    """Повертає ``(path, filename)`` за токеном або ``None``."""
    batch_dir = _batch_dir(batch_id)
    if not batch_dir:
        return None
    manifest_path = os.path.join(batch_dir, "manifest.json")
    if not os.path.isfile(manifest_path):
        return None
    with open(manifest_path, encoding="utf-8") as fh:
        manifest = json.load(fh)
    for f in manifest["files"]:
        if f["token"] == token and os.path.isfile(f["path"]):
            return f["path"], f["filename"]
    return None


def delete_batch(batch_id):
    batch_dir = _batch_dir(batch_id)
    if batch_dir:
        shutil.rmtree(batch_dir, ignore_errors=True)
        return True
    return False


def cleanup_stale_batches(max_age_hours=24):
    """Прибирає осиротілі теки (якщо Apps Script не викликав confirm)."""
    if not os.path.isdir(WZ_TMP_ROOT):
        return
    now = datetime.now().timestamp()
    for name in os.listdir(WZ_TMP_ROOT):
        path = os.path.join(WZ_TMP_ROOT, name)
        if os.path.isdir(path) and now - os.path.getmtime(path) > max_age_hours * 3600:
            shutil.rmtree(path, ignore_errors=True)
